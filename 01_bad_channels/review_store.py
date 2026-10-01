"""Canonical workbook for channel review; proposals and source EEG stay unchanged.
Fully written by an LLM
"""
from __future__ import annotations

import argparse
from contextlib import contextmanager
from datetime import datetime, timezone
import fcntl
import hashlib
import json
import os
from pathlib import Path
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
HEADERS = ("Channel name", "Flagged by", "Manual override", "Final decision", "Note")
FIRST_ROW = 6


def reasons(proposal, channel):
    labels = {"psd": "PSD", "ransac": "RANSAC", "SNR": "SNR", "hf_noise": "HF noise", "nan": "NaN"}
    keys = [k.removeprefix("bad_by_") for k, channels in proposal["bads"].items()
            if k.startswith("bad_by_") and channel in channels]
    return "; ".join(labels.get(k, k.replace("_", " ")) for k in keys) or "-"


def sheet_name(key):
    subject, run = key.split("_run-")
    return subject.removeprefix("sub-") + ("_OddGamble" if run == "01" else "_Axon")


class ReviewStore:
    """Read fresh decisions on each operation; save partial changes atomically."""

    def __init__(self, root=ROOT, path=None):
        self.root = Path(root).resolve()
        self.path = Path(path) if path else self.root / Path(__file__).resolve().parent.name / "channel_review.xlsx"
        self.recordings = {}
        self.digests = {}
        files = sorted((self.root / "logs/pyprep").glob("sub-*_run-*.json"))
        for file in files:
            data = file.read_bytes()
            proposal = json.loads(data)
            key = proposal["recording"]
            settings = proposal["settings"]
            if settings.get("random_state") != 2026 or settings.get("n_samples") != 1000:
                raise ValueError(f"{key}: finish the seed-2026 / 1000-subset detection run first")
            self.recordings[key] = proposal
            self.digests[key] = hashlib.sha256(data).hexdigest()
        expected = {f"sub-{s:03}_run-{r:02}" for s in range(1, 18) for r in (1, 2)}
        if set(self.recordings) != expected:
            raise ValueError("Need all 34 completed PyPREP recordings before creating/reviewing the workbook")

    @contextmanager
    def _lock(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self.path.with_suffix(".lock").open("a") as lock:
            fcntl.flock(lock.fileno(), fcntl.LOCK_EX)
            try:
                yield
            finally:
                fcntl.flock(lock.fileno(), fcntl.LOCK_UN)

    def _validate(self, workbook):
        if set(workbook.sheetnames) != {sheet_name(k) for k in self.recordings}:
            raise ValueError("Workbook sheets do not match the 34 recordings")
        for key, proposal in self.recordings.items():
            ws = workbook[sheet_name(key)]
            if ws["H1"].value != key or ws["H2"].value != self.digests[key]:
                raise ValueError(f"{key}: workbook belongs to different PyPREP proposals; preserve it and rebuild explicitly")
            if tuple(ws.cell(5, c).value for c in range(1, 6)) != HEADERS:
                raise ValueError(f"{key}: keep the spreadsheet column headings unchanged")
            if ws["C1"].value not in ("pending", "reviewed"):
                raise ValueError(f"{key}: review status must be pending or reviewed")
            names = proposal["eeg_ch_names"]
            if ws.max_row != FIRST_ROW + len(names) - 1:
                raise ValueError(f"{key}: unexpected channel rows")
            seen = set()
            for row in range(FIRST_ROW, FIRST_ROW + len(names)):
                ch = ws.cell(row, 1).value
                if ch not in names or ch in seen:
                    raise ValueError(f"{key}: missing, duplicate or unknown EEG channel {ch}")
                seen.add(ch)
                if ws.cell(row, 2).value != reasons(proposal, ch):
                    raise ValueError(f"{key}/{ch}: automatic proposal was edited; change Manual override instead")
                if ws.cell(row, 3).value not in ("yes", "-"):
                    raise ValueError(f"{key}/{ch}: Manual override must be yes or -")
                note = ws.cell(row, 5).value
                if note is not None and not isinstance(note, str):
                    raise ValueError(f"{key}/{ch}: note must be text")
        return workbook

    def _read(self):
        # Also guard an already-open GUI against a concurrently replaced proposal set.
        for key, expected in self.digests.items():
            proposal_path = self.root / f"logs/pyprep/{key}.json"
            if not proposal_path.exists() or hashlib.sha256(proposal_path.read_bytes()).hexdigest() != expected:
                raise ValueError(f"{key}: automatic proposals changed while the workbook was open; preserve decisions and migrate explicitly")
        if not self.path.exists():
            raise FileNotFoundError(f"Create the review workbook first: {self.path}")
        workbook = load_workbook(self.path)
        try:
            return self._validate(workbook)
        except Exception:
            workbook.close()
            raise

    def _write(self, workbook):
        with tempfile.NamedTemporaryFile(dir=self.path.parent, suffix=".xlsx", delete=False) as file:
            temp = Path(file.name)
        try:
            workbook.save(temp)
            os.replace(temp, self.path)
        finally:
            temp.unlink(missing_ok=True)

    def initialize(self):
        """Create once, translating existing final decisions into new overrides."""
        with self._lock():
            if self.path.exists():
                self._read().close()
                return 0
            legacy_path = self.root / "logs/review/decisions.json"
            legacy = json.loads(legacy_path.read_text()) if legacy_path.exists() else {}
            workbook = Workbook()
            workbook.remove(workbook.active)
            migrated = []
            for key, proposal in self.recordings.items():
                ws = workbook.create_sheet(sheet_name(key))
                ws["A1"] = sheet_name(key).replace("_", " · ")
                ws["B1"] = "Review status"
                ws["C1"] = "pending"
                ws["A2"] = "PyPREP seed 2026 · RANSAC n_samples = 1000"
                ws.merge_cells("A2:E2")
                ws["A3"] = "Manual override: yes reverses the automatic decision; - accepts it."
                ws.merge_cells("A3:E3")
                for col, label in enumerate(HEADERS, 1):
                    cell = ws.cell(5, col, label)
                    cell.fill = PatternFill("solid", fgColor="173B56")
                    cell.font = Font(color="FFFFFF", bold=True)
                for row, channel in enumerate(proposal["eeg_ch_names"], FIRST_ROW):
                    flagged = channel in proposal["all_bads"]
                    choice = legacy.get(key, {}).get(channel, {})
                    decision = choice.get("decision")
                    override = decision in ("keep", "exclude") and ((decision == "exclude") != flagged)
                    values = [channel, reasons(proposal, channel), "yes" if override else "-",
                              f'=IF(XOR(B{row}<>"-",C{row}="yes"),"bad","keep")', choice.get("note", "")]
                    for col, value in enumerate(values, 1):
                        ws.cell(row, col, value)
                    ws.cell(row, 5).data_type = "s"  # Notes are literal text, never spreadsheet formulas.
                    ws.row_dimensions[row].height = 32 if flagged or choice.get("note") else 21
                    for column in (3, 5):
                        ws.cell(row, column).fill = PatternFill("solid", fgColor="FFFCF0")
                    if decision in ("keep", "exclude"):
                        migrated.append(dict(recording=key, channel=channel, final_decision=decision,
                                             previous_updated_at=choice.get("updated_at")))
                for row, (label, value) in enumerate([
                    ("recording", key), ("proposal_sha256", self.digests[key]),
                    ("updated_utc", datetime.now(timezone.utc).isoformat()), ("schema", 1),
                    ("random_state", 2026), ("n_samples", 1000)], 1):
                    ws.cell(row, 7, label)
                    ws.cell(row, 8, value)
                ws.column_dimensions["G"].hidden = True
                ws.column_dimensions["H"].hidden = True
                for col, width in zip("ABCDE", (19, 38, 20, 18, 78)):
                    ws.column_dimensions[col].width = width
                ws.freeze_panes = "C6"
                last = FIRST_ROW + len(proposal["eeg_ch_names"]) - 1
                ws.auto_filter.ref = f"A5:E{last}"
                for column, formula, cells in [("override", '"-,yes"', f"C6:C{last}"),
                                                ("status", '"pending,reviewed"', "C1")]:
                    validation = DataValidation(type="list", formula1=formula, allow_blank=False)
                    validation.errorTitle = f"Invalid {column}"
                    validation.error = "Choose one of the listed values."
                    validation.showErrorMessage = True
                    validation.errorStyle = "stop"
                    ws.add_data_validation(validation)
                    validation.add(cells)
                ws.conditional_formatting.add(f"A6:E{last}", FormulaRule(formula=['$C6="yes"'],
                    fill=PatternFill("solid", fgColor="FFF1CC")))
                for row in ws.iter_rows(min_row=6, max_row=last, max_col=5):
                    for cell in row:
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
                ws.sheet_properties.pageSetUpPr.fitToPage = True
                ws.page_setup.orientation = "landscape"
                ws.page_setup.paperSize = ws.PAPERSIZE_A4
                ws.page_setup.fitToWidth = 1
                ws.page_setup.fitToHeight = 0
                ws.print_title_rows = "1:5"
                ws.print_area = f"A1:E{last}"
            self._validate(workbook)
            self._write(workbook)
            manifest = dict(created_utc=datetime.now(timezone.utc).isoformat(), random_state=2026,
                            n_samples=1000, recordings=34, proposal_sha256=self.digests,
                            migrated_decisions=migrated, review_status="All recordings start pending for the new proposals.")
            (self.path.parent / "review_provenance.json").write_text(json.dumps(manifest, indent=2) + "\n")
            return len(migrated)

    def reload(self):
        workbook = self._read()
        workbook.close()

    def choices(self, key):
        workbook = self._read()
        try:
            ws = workbook[sheet_name(key)]
            flagged = set(self.recordings[key]["all_bads"])
            return {ws.cell(row, 1).value: dict(
                        override=ws.cell(row, 3).value == "yes", note=ws.cell(row, 5).value or "",
                        final_bad=(ws.cell(row, 1).value in flagged) != (ws.cell(row, 3).value == "yes"))
                    for row in range(FIRST_ROW, FIRST_ROW + len(self.recordings[key]["eeg_ch_names"]))}
        finally:
            workbook.close()

    def is_reviewed(self, key):
        workbook = self._read()
        try:
            return workbook[sheet_name(key)]["C1"].value == "reviewed"
        finally:
            workbook.close()

    def save(self, key, overrides, notes=None, reviewed=False):
        if key not in self.recordings:
            raise KeyError(key)
        notes = notes or {}
        names = set(self.recordings[key]["eeg_ch_names"])
        if not (set(overrides) | set(notes)) <= names:
            raise ValueError("Only recorded EEG channels can receive an override")
        if any(type(value) is not bool for value in overrides.values()):
            raise ValueError("Overrides must be booleans")
        if any(not isinstance(value, str) for value in notes.values()):
            raise ValueError("Notes must be strings")
        with self._lock():
            workbook = self._read()
            try:
                ws = workbook[sheet_name(key)]
                for row in range(FIRST_ROW, FIRST_ROW + len(names)):
                    ch = ws.cell(row, 1).value
                    if ch in overrides:
                        ws.cell(row, 3, "yes" if overrides[ch] else "-")
                    if ch in notes:
                        ws.cell(row, 5, notes[ch])
                        ws.cell(row, 5).data_type = "s"
                    # Refresh derived formulas after an external spreadsheet sort.
                    ws.cell(row, 4, f'=IF(XOR(B{row}<>"-",C{row}="yes"),"bad","keep")')
                ws["C1"] = "reviewed" if reviewed else "pending"
                ws["H3"] = datetime.now(timezone.utc).isoformat()
                self._write(workbook)
            finally:
                workbook.close()

    def export_decisions(self, require_reviewed=True):
        workbook = self._read()
        try:
            pending = [key for key in self.recordings if workbook[sheet_name(key)]["C1"].value != "reviewed"]
            if require_reviewed and pending:
                raise ValueError(f"{len(pending)} recordings still pending review: " + ", ".join(pending[:6]))
            output = {}
            for key, proposal in self.recordings.items():
                ws = workbook[sheet_name(key)]
                flagged = set(proposal["all_bads"])
                output[key] = {}
                for row in range(FIRST_ROW, FIRST_ROW + len(proposal["eeg_ch_names"])):
                    ch = ws.cell(row, 1).value
                    bad = (ch in flagged) != (ws.cell(row, 3).value == "yes")
                    output[key][ch] = dict(decision="exclude" if bad else "keep",
                                          note=ws.cell(row, 5).value or "", updated_at=ws["H3"].value,
                                          automatic_bad=ch in flagged, manual_override=ws.cell(row, 3).value == "yes")
            return output
        finally:
            workbook.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=ROOT)
    args = parser.parse_args()
    store = ReviewStore(args.root)
    count = store.initialize()
    print(f"Workbook ready: {store.path}; {count} earlier decisions imported on creation.")
